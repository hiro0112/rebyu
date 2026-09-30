"""Amazon 商品ページのレビューを取得し、点数別に分けて Excel に出力する。

使い方:
    python analyze_reviews.py <ASIN> [保存済みHTMLのパス]

注意:
    ログインなしで取得できるのは商品ページに表示されるレビュー（通常 8 件前後）のみ。
    全件のレビュー一覧ページ (/product-reviews/) はログイン画面へリダイレクトされる。
"""
import re
import sys
import urllib.request
from collections import Counter
from datetime import date
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36")

# 話題タグ: 本文・タイトルにキーワードが含まれていれば付与する
TOPICS = {
    "作りやすさ": ["簡単", "作りやす", "見やす", "チャレンジしやす", "楽しく", "綺麗に出来", "見慣れ"],
    "デザイン・見た目": ["デザイン", "可愛", "モチーフ", "絵", "写真通り", "素敵", "華やか"],
    "糸・材料": ["糸", "材料", "生地"],
    "付属の枠": ["額縁", "木枠"],
    "サイズ": ["サイズ", "大き"],
    "季節・贈り物": ["お正月", "年末", "祖母", "あげた"],
    "不満・要望": ["足りな", "安っぽ", "見にく", "よかったか微妙", "選ばなかった"],
}

# 点数別の傾向分析（取得したレビュー本文を読んだうえでの所見）
ANALYSIS = {
    5: {
        "傾向": "作りやすさ・デザインの良さ・季節に合わせて完成できた満足感が中心。",
        "詳細": [
            "「簡単」「作りやすい」「説明書が見やすい」など、初心者〜経験者が迷わず完成できた点を評価する声が最多（4/5件）。",
            "朝顔・椿など季節モチーフのデザインや、微妙に色の違う糸が多く入っている点を「可愛い」「ワクワクする」と評価。",
            "お正月に間に合わせた、祖母が喜んだなど、季節行事や家族への贈り物としての満足感が語られている。",
            "高評価でも「糸が少し足りなかった」「木枠がやや安っぽい」「最初は図案が見にくい」という軽い不満が併記されている。",
        ],
    },
    4: {
        "傾向": "品質には満足しているが、サイズや華やかさなど『期待とのギャップ』で1点減。",
        "詳細": [
            "「写真通り」「材料も十分」と品物自体への不満はない。",
            "減点理由は「もう少し大きい方がよかった」（約17×17cmのサイズ感）、「華やかさを求めるなら選ばなかった」（落ち着いた色合い）。",
            "いずれも不良・欠品ではなく、商品ページから完成サイズや印象が伝わりきっていないことによるギャップ。",
        ],
    },
    3: {"傾向": "本文付きレビューなし（評価のみ）。", "詳細": ["ログインなしでは本文を確認できないため、傾向は分析不可。"]},
    2: {"傾向": "本文付きレビューなし（評価のみ）。", "詳細": ["ログインなしでは本文を確認できないため、傾向は分析不可。"]},
    1: {"傾向": "該当なし（評価 0%）。", "詳細": ["星1の評価は付いていない。"]},
}

SUMMARY = [
    "平均 4.4 / 5、星4以上が約89%と満足度は高い。星1は0%。",
    "高評価の最大の理由は『作りやすさ』。図案・説明書が分かりやすく、クロスステッチ初心者でも完成できるという声が多い。",
    "季節モチーフ（お正月・朝顔など）が行事や贈り物の用途と合っていて、完成後の満足度につながっている。",
    "不満点は『サイズが小さい』『華やかさが控えめ』『糸がやや不足』『木枠が安っぽい』。致命的なものはなく、期待値の調整で防げる内容が中心。",
    "改善のヒント: 商品ページで完成サイズ（約17cm角）の実物イメージや色味を明示し、糸の予備を少し増やすと星4→星5への引き上げが見込める。",
    "※ 取得できた本文付きレビューは8件（星5×5、星4×3）のみ。星3・星2の理由は未確認のため、結論は参考値として扱うこと。",
    "※ Amazon ではパターン違い（バラ・朝顔・マーガレット等）のレビューが共通表示されるため、『椿のお正月飾り』自体のレビューは3件。",
]

HEADERS = ["No", "評価", "タイトル", "本文", "パターン", "投稿日", "投稿者", "購入確認", "参考になった数", "話題タグ"]
WIDTHS = [5, 7, 26, 60, 18, 14, 16, 10, 12, 26]

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF")
STAR_FILLS = {5: "E2EFDA", 4: "EEF5E6", 3: "FFF2CC", 2: "FCE4D6", 1: "F8CBAD"}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def fetch(asin: str) -> str:
    req = urllib.request.Request(
        f"https://www.amazon.co.jp/dp/{asin}",
        headers={"User-Agent": UA, "Accept-Language": "ja-JP,ja;q=0.9"},
    )
    with urllib.request.urlopen(req, timeout=30) as res:
        return res.read().decode("utf-8", errors="replace")


def parse(html: str):
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True).replace("Amazon.co.jp: ", "")
    total = soup.select_one('[data-hook="total-review-count"]')
    total = int(re.sub(r"\D", "", total.get_text())) if total else 0
    avg = soup.select_one('[data-hook="rating-out-of-text"]')
    avg = float(re.search(r"([\d.]+)つ$", avg.get_text()).group(1)) if avg else None
    dist = {}
    # 各 li の末尾の % が、その星（5→1 の順）の割合
    for star, li in zip(range(5, 0, -1), soup.select("#histogramTable li")):
        m = re.search(r"(\d+)%$", li.get_text(" ", strip=True))
        if m:
            dist[star] = int(m.group(1))

    reviews = []
    for r in soup.select('[data-hook="review"]'):
        def text(hook):
            e = r.select_one(f'[data-hook="{hook}"]')
            return e.get_text("\n", strip=True) if e else ""

        helpful = re.search(r"(\d+)", text("helpful-vote-statement"))
        reviews.append({
            "star": int(re.search(r"中(\d)", text("review-star-rating")).group(1)),
            "title": text("reviewTitle"),
            "body": text("reviewRichContentContainer"),
            "pattern": text("format-strip").replace("パターン名: ", ""),
            "date": re.sub(r"に.*", "", text("review-date")),
            "name": (r.select_one(".a-profile-name") or soup.new_tag("x")).get_text(strip=True),
            "verified": "Amazonで購入" in text("review-badges"),
            "helpful": int(helpful.group(1)) if helpful else 0,
        })
    for rv in reviews:
        s = rv["title"] + rv["body"]
        rv["tags"] = [t for t, kws in TOPICS.items() if any(k in s for k in kws)]
    return title, total, avg, dist, reviews


def style_header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_reviews(ws, reviews, empty_note=None):
    ws.append(HEADERS)
    style_header(ws)
    for i, w in enumerate(WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    for i, rv in enumerate(reviews, 1):
        ws.append([i, "★" * rv["star"], rv["title"], rv["body"], rv["pattern"], rv["date"], rv["name"],
                   "○" if rv["verified"] else "", rv["helpful"], "、".join(rv["tags"])])
        for c in ws[ws.max_row]:
            c.alignment, c.border = WRAP, BORDER
            c.fill = PatternFill("solid", fgColor=STAR_FILLS[rv["star"]])
    if not reviews and empty_note:
        ws.append([empty_note])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
        ws["A2"].font = Font(italic=True, color="808080")
    ws.freeze_panes = "A2"
    if reviews:
        ws.auto_filter.ref = ws.dimensions


def build(asin, title, total, avg, dist, reviews, out: Path):
    wb = Workbook()
    by_star = {s: [r for r in reviews if r["star"] == s] for s in range(5, 0, -1)}

    # 概要
    ws = wb.active
    ws.title = "概要"
    ws["A1"] = "Amazon レビュー分析"
    ws["A1"].font = Font(bold=True, size=14)
    info = [("商品", title), ("ASIN", asin), ("URL", f"https://www.amazon.co.jp/dp/{asin}"),
            ("取得日", date.today().isoformat()), ("平均評価", avg), ("総評価数", total),
            ("本文付きで取得できたレビュー", len(reviews))]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    r0 = 3 + len(info) + 1
    ws.cell(r0, 1, "点数")
    ws.cell(r0, 2, "割合")
    ws.cell(r0, 3, "推定評価数")
    ws.cell(r0, 4, "取得レビュー数")
    style_header(ws, r0)
    for i, s in enumerate(range(5, 0, -1), r0 + 1):
        pct = dist.get(s, 0)
        ws.cell(i, 1, f"星{s}")
        ws.cell(i, 2, pct / 100).number_format = "0%"
        ws.cell(i, 3, round(total * pct / 100))
        ws.cell(i, 4, len(by_star[s]))
        for c in ws[i][:4]:
            c.border = BORDER
    note_row = r0 + 7
    ws.cell(note_row, 1, "注意").font = Font(bold=True, color="C00000")
    ws.cell(note_row, 2, "Amazon のレビュー一覧ページはログインが必要なため、商品ページに表示される本文付きレビューのみを収集しています。"
                         "総評価数の多くは本文なしの評価のみです。推定評価数は総評価数×割合の概算です。")
    ws.cell(note_row, 2).alignment = WRAP
    ws.row_dimensions[note_row].height = 45
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 14

    chart = BarChart()
    chart.type = "bar"
    chart.title = "点数別の評価割合"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=r0, max_row=r0 + 5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + 5))
    chart.y_axis.numFmt = "0%"
    chart.y_axis.majorGridlines = None
    chart.x_axis.scaling.orientation = "maxMin"
    chart.height, chart.width = 7, 14
    ws.add_chart(chart, f"A{note_row + 2}")

    # 点数別シート
    for s, rows in by_star.items():
        est = round(total * dist.get(s, 0) / 100)
        note = f"本文付きレビューなし（評価のみ 約{est}件）" if est else "該当なし（評価 0件）"
        write_reviews(wb.create_sheet(f"星{s}"), rows, note)

    write_reviews(wb.create_sheet("全レビュー"), sorted(reviews, key=lambda r: -r["star"]))

    # 傾向分析
    ws = wb.create_sheet("傾向分析")
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 95
    ws["A1"] = "点数別の傾向"
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["点数", "取得件数", "傾向・詳細"])
    style_header(ws, 2)
    for s in range(5, 0, -1):
        a = ANALYSIS[s]
        body = "【" + a["傾向"] + "】\n" + "\n".join("・" + d for d in a["詳細"])
        ws.append([f"星{s}", len(by_star[s]), body])
        for c in ws[ws.max_row]:
            c.alignment, c.border = WRAP, BORDER
            c.fill = PatternFill("solid", fgColor=STAR_FILLS[s])
        ws.row_dimensions[ws.max_row].height = 18 * (len(a["詳細"]) + 1) + 8

    ws.append([])
    start = ws.max_row + 1
    ws.cell(start, 1, "話題別の言及数（点数別）").font = Font(bold=True, size=13)
    ws.append(["話題"] + [f"星{s}" for s in range(5, 0, -1)] + ["合計"])
    style_header(ws, start + 1)
    for t in TOPICS:
        cnt = Counter(r["star"] for r in reviews if t in r["tags"])
        ws.append([t] + [cnt.get(s, 0) for s in range(5, 0, -1)] + [sum(cnt.values())])
        for c in ws[ws.max_row][:7]:
            c.border = BORDER
    for col in "DEFG":
        ws.column_dimensions[col].width = 8
    ws.column_dimensions["A"].width = 16

    ws.append([])
    ws.cell(ws.max_row + 1, 1, "総合所見").font = Font(bold=True, size=13)
    for line in SUMMARY:
        ws.append(["", "", "・" + line])
        ws.cell(ws.max_row, 3).alignment = WRAP

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main():
    asin = sys.argv[1] if len(sys.argv) > 1 else "B0BFGLYJRN"
    html = Path(sys.argv[2]).read_text(encoding="utf-8") if len(sys.argv) > 2 else fetch(asin)
    title, total, avg, dist, reviews = parse(html)
    out = Path(__file__).parent / "output" / f"レビュー分析_{asin}.xlsx"
    build(asin, title, total, avg, dist, reviews, out)
    print(f"{len(reviews)} 件のレビューを出力: {out}")


if __name__ == "__main__":
    main()
